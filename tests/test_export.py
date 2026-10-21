"""Tests for core/export.py (design.md section 21, CA-36 and RF-E4)."""

from __future__ import annotations

from io import BytesIO

import numpy as np
from openpyxl import load_workbook

from core.export import MONTE_CARLO_ROW_CAP, ExportPayload, build_workbook
from core.optimize import check_feasibility, equal_weights
from core.stats import summarize_portfolio
from core.types import Constraints
from ui.labels import EXPORT_LABELS, SHEET_NAMES


def _workbook(payload: ExportPayload):
    blob = build_workbook(payload, sheet_names=SHEET_NAMES,
                          column_labels=EXPORT_LABELS)
    assert isinstance(blob, bytes) and len(blob) > 0
    return load_workbook(BytesIO(blob))


def _equal_weight_optimization(stats):
    weights = equal_weights(len(stats.names))
    constraints = Constraints(allow_short=False, max_weight=1.0)
    summary = summarize_portfolio(
        'equal_weight', weights, stats.mu_annual, stats.cov_annual,
        stats.convention.rf_cc_annual, converged=True,
        feasibility=check_feasibility(weights, constraints),
        diagnostics={'message': ''})
    return {'equal_weight': summary}


def test_workbook_is_generated_and_reopens(stats):
    payload = ExportPayload(params={'rf_annual_effective': 0.045, 'seed': 42},
                            stats=stats,
                            optimization=_equal_weight_optimization(stats))
    book = _workbook(payload)
    for key in ('parameters', 'summary', 'weights', 'asset_stats', 'covariance',
                'correlation'):
        assert SHEET_NAMES[key] in book.sheetnames
    assert book.sheetnames[0] == SHEET_NAMES['parameters']
    sheet = book[SHEET_NAMES['asset_stats']]
    assert sheet.freeze_panes == 'A2'
    assert sheet.max_row == len(stats.names) + 1
    assert sheet['A1'].font.bold is True


def test_sheet_titles_fit_the_excel_limit():
    assert all(len(title) <= 31 for title in SHEET_NAMES.values())


def test_missing_sections_are_skipped_and_annotated():
    payload = ExportPayload(params={'seed': 42})
    book = _workbook(payload)
    assert book.sheetnames == [SHEET_NAMES['parameters']]
    sheet = book[SHEET_NAMES['parameters']]
    texts = [str(cell.value) for row in sheet.iter_rows() for cell in row
             if cell.value is not None]
    assert EXPORT_LABELS['missing_sheets'] in texts
    assert texts.count(EXPORT_LABELS['not_computed']) >= 10
    for key in ('summary', 'backtest', 'monte_carlo_capacity'):
        assert EXPORT_LABELS.get(key, key) in texts or key in texts


def test_monte_carlo_weights_is_truncated_with_a_note():
    total = MONTE_CARLO_ROW_CAP + 5
    cloud = (np.linspace(0.1, 0.4, total), np.linspace(0.0, 0.2, total),
             np.linspace(-0.5, 2.0, total))
    book = _workbook(ExportPayload(params={}, weight_cloud=cloud))
    sheet = book[SHEET_NAMES['monte_carlo_weights']]
    # header + cap rows + blank row + note row
    assert sheet.max_row == MONTE_CARLO_ROW_CAP + 3
    note = str(sheet.cell(row=sheet.max_row, column=1).value)
    assert EXPORT_LABELS['monte_carlo_rows_truncated'] in note
    assert f'{MONTE_CARLO_ROW_CAP}/{total}' in note


def test_a_broken_sheet_does_not_break_the_download():
    class Exploding:
        @property
        def rows(self):
            raise RuntimeError('boom')

        def __getattr__(self, item):
            raise RuntimeError('boom')

    book = _workbook(ExportPayload(params={}, purchase=Exploding()))
    sheet = book[SHEET_NAMES['purchase_balance']]
    assert EXPORT_LABELS['sheet_error'] in str(sheet.cell(row=2, column=1).value)


def test_percent_and_money_formats_are_applied(stats):
    payload = ExportPayload(params={}, stats=stats,
                            optimization=_equal_weight_optimization(stats))
    book = _workbook(payload)
    sheet = book[SHEET_NAMES['weights']]
    assert sheet.cell(row=2, column=2).number_format == '0.00%'
    stats_sheet = book[SHEET_NAMES['asset_stats']]
    assert stats_sheet.cell(row=2, column=3).number_format == '0.00%'
